"""Buat templat Excel raport TPQ gratis (SEO Tahap 3, 26 Sep 2026).

Keluaran: static/unduh/templat-raport-tpq.xlsx
Jalankan:  python3 scripts/buat-templat-raport-tpq.py   (butuh openpyxl)

Isi:
  Petunjuk  — cara pakai, arti warna, ambang predikat (bisa diubah TPQ)
  Identitas — data TPQ & semester (dipakai lembar Rapor)
  Nilai     — satu baris per santri (40 baris), rata-rata & predikat otomatis
  Rapor     — pilih nomor santri → rapor siap cetak A4

Hanya fungsi Excel 2007 (IF, AVERAGE, COUNT, INDEX, MATCH, IFERROR) supaya
jalan di Excel lama, LibreOffice, WPS, dan Google Sheets.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

KELUAR = Path(__file__).resolve().parent.parent / 'static' / 'unduh' / 'templat-raport-tpq.xlsx'
FONT = 'Arial'
HIJAU = '1F5B3F'
KUNING = 'FFF2B3'  # sel yang diisi guru
KRIM = 'F7F3E8'
tipis = Side(style='thin', color='B8B0A0')
KOTAK = Border(left=tipis, right=tipis, top=tipis, bottom=tipis)
ISI = PatternFill('solid', fgColor=KUNING)
KEPALA = PatternFill('solid', fgColor=HIJAU)
LATAR = PatternFill('solid', fgColor=KRIM)


def f(size=10, bold=False, color='000000', italic=False):
	return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


wb = Workbook()

# ------------------------------------------------------------------ Petunjuk
p = wb.active
p.title = 'Petunjuk'
p.column_dimensions['A'].width = 3
p.column_dimensions['B'].width = 30
p.column_dimensions['C'].width = 14
p.column_dimensions['D'].width = 60
p['B2'] = 'Templat Raport TPQ — gratis dari SantriOnline'
p['B2'].font = f(16, True, HIJAU)
p['B3'] = 'Boleh dipakai, diubah, dan dibagikan untuk TPQ mana pun. Tanpa makro, tanpa kata sandi.'
p['B3'].font = f(10, italic=True)
langkah = [
	'1. Isi lembar "Identitas": nama TPQ, semester, tahun ajaran, nama kepala TPQ dan wali kelas.',
	'2. Isi lembar "Nilai": satu baris satu santri. Nilai 0–100. Rata-rata dan predikat terisi sendiri.',
	'3. Buka lembar "Rapor", ketik nomor urut santri di sel kuning C3, lalu cetak (A4, tegak).',
	'4. Ulangi langkah 3 untuk santri berikutnya. Tidak perlu mengetik ulang apa pun.',
	'Baris contoh (No. 1, "Ahmad Fauzan") boleh dihapus atau ditimpa.',
]
for i, t in enumerate(langkah):
	p.cell(row=5 + i, column=2, value=t).font = f(10)
p['B11'] = 'Arti warna'
p['B11'].font = f(11, True, HIJAU)
p['B12'] = 'Sel kuning'
p['B12'].fill = ISI
p['B12'].border = KOTAK
p['D12'] = 'Diisi guru/pengurus.'
p['B13'] = 'Sel putih/hijau'
p['B13'].border = KOTAK
p['D13'] = 'Rumus otomatis — jangan diketik manual.'
for r in (12, 13):
	p.cell(row=r, column=2).font = f(10)
	p.cell(row=r, column=4).font = f(10)

p['B15'] = 'Ambang predikat (boleh diubah sesuai kebijakan TPQ)'
p['B15'].font = f(11, True, HIJAU)
p['B16'], p['C16'], p['D16'] = 'Predikat', 'Nilai minimal', 'Keterangan'
for c in 'BCD':
	p[f'{c}16'].font = f(10, True, 'FFFFFF')
	p[f'{c}16'].fill = KEPALA
AMBANG = [
	('Mumtaz', 90, 'Istimewa — lancar, hampir tanpa kesalahan.'),
	('Jayyid Jiddan', 80, 'Sangat baik.'),
	('Jayyid', 70, 'Baik.'),
	('Maqbul', 60, 'Cukup — perlu penguatan.'),
	('Perlu bimbingan', 0, 'Belum mencapai batas cukup; ulangi bersama guru.'),
]
for i, (nama, batas, ket) in enumerate(AMBANG):
	r = 17 + i
	p.cell(row=r, column=2, value=nama).font = f(10, True)
	c = p.cell(row=r, column=3, value=batas)
	c.font = f(10, color='0000FF')
	c.fill = ISI
	p.cell(row=r, column=4, value=ket).font = f(10)
	for col in (2, 3, 4):
		p.cell(row=r, column=col).border = KOTAK
p['C17'].comment = Comment(
	'Ambang umum yang banyak dipakai lembaga tahfidz di Indonesia (Mumtaz 90–100, Jayyid Jiddan 80–89, Jayyid 70–79). '
	'Bukan standar resmi Kemenag — sesuaikan dengan pedoman TPQ Anda.',
	'SantriOnline',
)
p['B23'] = 'Ingin rapor tersusun sendiri dari catatan setoran harian, dan bisa dilihat wali dari HP?'
p['B23'].font = f(10, True, HIJAU)
p['B24'] = 'Coba aplikasi SantriOnline untuk TPQ (gratis): https://santrionline.com/aplikasi-raport-tpq'
p['B24'].font = f(10, color='0563C1')
p['B24'].hyperlink = 'https://santrionline.com/aplikasi-raport-tpq'

# ----------------------------------------------------------------- Identitas
d = wb.create_sheet('Identitas')
d.column_dimensions['A'].width = 3
d.column_dimensions['B'].width = 26
d.column_dimensions['C'].width = 44
d['B2'] = 'Identitas TPQ & Semester'
d['B2'].font = f(14, True, HIJAU)
IDENT = [
	('Nama TPQ', 'TPQ Al-Hidayah'),
	('Alamat', 'Jl. Masjid No. 1, Desa ...'),
	('Semester', 'Ganjil'),
	('Tahun ajaran', '2026/2027'),
	('Kepala TPQ', 'Ust. ...'),
	('Wali kelas / ustadz', 'Ust. ...'),
	('Tempat & tanggal rapor', 'Batu, 20 Desember 2026'),
]
for i, (k, v) in enumerate(IDENT):
	r = 4 + i
	d.cell(row=r, column=2, value=k).font = f(10, True)
	c = d.cell(row=r, column=3, value=v)
	c.font = f(10, color='0000FF')
	c.fill = ISI
	c.border = KOTAK
dv_sem = DataValidation(type='list', formula1='"Ganjil,Genap"', allow_blank=True)
d.add_data_validation(dv_sem)
dv_sem.add('C6')

# --------------------------------------------------------------------- Nilai
n = wb.create_sheet('Nilai')
KOLOM = [
	('No', 5), ('Nama santri', 24), ('NIS', 10), ('Jilid / kelas', 12),
	('Kelancaran membaca (tartil)', 12), ('Tajwid', 10), ('Makharijul huruf', 11),
	('Hafalan surat pendek', 11), ('Hafalan doa harian', 11), ('Praktik wudhu & shalat', 11),
	('Menulis huruf Arab', 11), ('Rata-rata', 10), ('Predikat', 15),
	('Sakit', 7), ('Izin', 7), ('Alpa', 7), ('Adab (A/B/C)', 9), ('Catatan ustadz', 40),
]
NILAI_COL = list(range(5, 12))  # E..K
n['A1'] = 'Nilai santri — isi sel kuning; kolom Rata-rata & Predikat otomatis'
n['A1'].font = f(12, True, HIJAU)
for i, (judul, lebar) in enumerate(KOLOM, start=1):
	c = n.cell(row=3, column=i, value=judul)
	c.font = f(9, True, 'FFFFFF')
	c.fill = KEPALA
	c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
	c.border = KOTAK
	n.column_dimensions[c.column_letter].width = lebar
n.row_dimensions[3].height = 44
BARIS_AWAL, JUMLAH = 4, 40
BARIS_AKHIR = BARIS_AWAL + JUMLAH - 1
contoh = [1, 'Ahmad Fauzan', '2026-001', 'Jilid 4', 85, 80, 82, 90, 88, 86, 78, None, None, 1, 0, 0, 'A',
		  'Bacaan sudah lancar; perhatikan panjang mad thabi\'i.']
for r in range(BARIS_AWAL, BARIS_AKHIR + 1):
	no = n.cell(row=r, column=1, value=r - BARIS_AWAL + 1)
	no.font = f(10)
	no.alignment = Alignment(horizontal='center')
	for col in range(2, 19):
		c = n.cell(row=r, column=col)
		c.border = KOTAK
		c.font = f(10)
		if col not in (12, 13):
			c.fill = ISI
			c.font = f(10, color='0000FF')
	no.border = KOTAK
	rng = f'E{r}:K{r}'
	n.cell(row=r, column=12, value=f'=IF(COUNT({rng})=0,"",ROUND(AVERAGE({rng}),1))')
	n.cell(row=r, column=13, value=(
		f'=IF(L{r}="","",IF(L{r}>=Petunjuk!$C$17,Petunjuk!$B$17,IF(L{r}>=Petunjuk!$C$18,Petunjuk!$B$18,'
		f'IF(L{r}>=Petunjuk!$C$19,Petunjuk!$B$19,IF(L{r}>=Petunjuk!$C$20,Petunjuk!$B$20,Petunjuk!$B$21)))))'
	))
	n.cell(row=r, column=18).alignment = Alignment(wrap_text=True, vertical='top')
for col, v in enumerate(contoh, start=1):
	if col in (1, 12, 13) or v is None:
		continue
	n.cell(row=BARIS_AWAL, column=col, value=v)
n.freeze_panes = 'C4'
dv_nilai = DataValidation(type='decimal', operator='between', formula1='0', formula2='100', allow_blank=True,
						  error='Nilai 0–100', errorTitle='Nilai tidak valid')
dv_adab = DataValidation(type='list', formula1='"A,B,C"', allow_blank=True)
n.add_data_validation(dv_nilai)
n.add_data_validation(dv_adab)
dv_nilai.add(f'E{BARIS_AWAL}:K{BARIS_AKHIR}')
dv_adab.add(f'Q{BARIS_AWAL}:Q{BARIS_AKHIR}')

# --------------------------------------------------------------------- Rapor
rp = wb.create_sheet('Rapor')
rp.sheet_properties.pageSetUpPr.fitToPage = True
rp.page_setup.orientation = 'portrait'
rp.page_setup.paperSize = rp.PAPERSIZE_A4
rp.page_setup.fitToWidth = 1
rp.page_setup.fitToHeight = 1
rp.print_options.horizontalCentered = True
for col, w in zip('ABCDE', (3, 34, 14, 18, 22)):
	rp.column_dimensions[col].width = w
rp['B3'] = 'Nomor urut santri →'
rp['B3'].font = f(10, True)
rp['B3'].alignment = Alignment(horizontal='right')
rp['C3'] = 1
rp['C3'].fill = ISI
rp['C3'].font = f(12, True, '0000FF')
rp['C3'].border = KOTAK
rp['D3'] = '(ketik nomor dari lembar Nilai, lalu cetak)'
rp['D3'].font = f(8, italic=True, color='777777')
NILAI = f"Nilai!$A${BARIS_AWAL}:$R${BARIS_AKHIR}"
BARIS_SANTRI = f'MATCH($C$3,Nilai!$A${BARIS_AWAL}:$A${BARIS_AKHIR},0)'


def ambil(kolom: int) -> str:
	return f'=IFERROR(INDEX({NILAI},{BARIS_SANTRI},{kolom})&"","")'


def ambil_angka(kolom: int) -> str:
	return f'=IFERROR(IF(INDEX({NILAI},{BARIS_SANTRI},{kolom})="","",INDEX({NILAI},{BARIS_SANTRI},{kolom})),"")'


rp['B5'] = '=UPPER(Identitas!$C$4)'
rp['B5'].font = f(15, True, HIJAU)
rp['B6'] = '=Identitas!$C$5'
rp['B6'].font = f(9, color='555555')
rp['B7'] = 'LAPORAN HASIL BELAJAR SANTRI'
rp['B7'].font = f(12, True)
for cel in ('B5', 'B6', 'B7'):
	rp.merge_cells(f'{cel}:{cel[0].replace("B", "E")}{cel[1:]}')
	rp[cel].alignment = Alignment(horizontal='center')
IDN = [
	(9, 'Nama santri', ambil(2)), (10, 'NIS', ambil(3)), (11, 'Jilid / kelas', ambil(4)),
	(12, 'Semester / tahun ajaran', '=Identitas!$C$6&" / "&Identitas!$C$7'),
]
for r, k, v in IDN:
	rp.cell(row=r, column=2, value=k).font = f(10)
	c = rp.cell(row=r, column=3, value=v)
	c.font = f(10, True)
	rp.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
for i, judul in enumerate(('Aspek yang dinilai', 'Nilai', 'Predikat'), start=2):
	c = rp.cell(row=14, column=i, value=judul)
	c.font = f(10, True, 'FFFFFF')
	c.fill = KEPALA
	c.border = KOTAK
	c.alignment = Alignment(horizontal='center')
rp.merge_cells('D14:E14')
PRED = ('=IF(C{r}="","",IF(C{r}>=Petunjuk!$C$17,Petunjuk!$B$17,IF(C{r}>=Petunjuk!$C$18,Petunjuk!$B$18,'
		'IF(C{r}>=Petunjuk!$C$19,Petunjuk!$B$19,IF(C{r}>=Petunjuk!$C$20,Petunjuk!$B$20,Petunjuk!$B$21)))))')
for i, col in enumerate(NILAI_COL):
	r = 15 + i
	rp.cell(row=r, column=2, value=KOLOM[col - 1][0]).font = f(10)
	rp.cell(row=r, column=3, value=ambil_angka(col)).font = f(10)
	rp.cell(row=r, column=4, value=PRED.format(r=r)).font = f(10)
	rp.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
	for col2 in (2, 3, 4, 5):
		rp.cell(row=r, column=col2).border = KOTAK
	rp.cell(row=r, column=3).alignment = Alignment(horizontal='center')
r = 15 + len(NILAI_COL)  # 22
rp.cell(row=r, column=2, value='Rata-rata').font = f(10, True)
rp.cell(row=r, column=3, value=ambil_angka(12)).font = f(10, True)
rp.cell(row=r, column=4, value=ambil(13)).font = f(10, True)
rp.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
for col2 in (2, 3, 4, 5):
	rp.cell(row=r, column=col2).border = KOTAK
	rp.cell(row=r, column=col2).fill = LATAR
rp.cell(row=r, column=3).alignment = Alignment(horizontal='center')

rp['B24'] = 'Kehadiran & adab'
rp['B24'].font = f(10, True, HIJAU)
KH = [(25, 'Sakit (hari)', 14), (26, 'Izin (hari)', 15), (27, 'Alpa (hari)', 16), (28, 'Adab', 17)]
for rr, k, col in KH:
	rp.cell(row=rr, column=2, value=k).font = f(10)
	rp.cell(row=rr, column=3, value=ambil_angka(col) if col != 17 else ambil(col)).font = f(10)
	for col2 in (2, 3):
		rp.cell(row=rr, column=col2).border = KOTAK
	rp.cell(row=rr, column=3).alignment = Alignment(horizontal='center')
rp['B30'] = 'Catatan ustadz'
rp['B30'].font = f(10, True, HIJAU)
rp['B31'] = ambil(18)
rp['B31'].font = f(10)
rp['B31'].alignment = Alignment(wrap_text=True, vertical='top')
rp.merge_cells('B31:E33')
rp.row_dimensions[31].height = 24
for rr in range(31, 34):
	for col2 in range(2, 6):
		rp.cell(row=rr, column=col2).border = KOTAK

rp['E35'] = '=Identitas!$C$10'
rp['B36'] = 'Mengetahui, Kepala TPQ'
rp['E36'] = 'Wali kelas / ustadz'
rp['B37'] = 'Orang tua / wali santri'
rp['B40'] = '=Identitas!$C$8'
rp['E40'] = '=Identitas!$C$9'
rp['B37'].value = None
for cel in ('E35', 'B36', 'E36', 'B40', 'E40'):
	rp[cel].font = f(10, cel in ('B40', 'E40'))
rp['C36'] = 'Orang tua / wali'
rp['C40'] = '(....................)'
rp['C36'].font = f(10)
rp['C40'].font = f(10)
rp.print_area = 'A5:E42'
rp['B42'] = 'Templat gratis dari santrionline.com — rapor otomatis dari setoran harian: santrionline.com/aplikasi-raport-tpq'
rp['B42'].font = f(7, italic=True, color='888888')

wb.active = wb.sheetnames.index('Petunjuk')
KELUAR.parent.mkdir(parents=True, exist_ok=True)
wb.save(KELUAR)
print(KELUAR)
